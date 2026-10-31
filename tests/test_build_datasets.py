import os
import sys
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import datetime

import fitz
import openpyxl

from build_datasets import (RAW_DIR, VHTS_TPK_CSV, hotel_formula_diffs, hotel_record_from_row, jampea_records_from_rows,
                            jampea_rows_from_sheet, jampea_total_diffs, normalize_komoditas, parse_id_number,
                            parse_jampea_pdf, parse_jampea_words_pdf, parse_period_from_path, read_vhts_tpk)


class ParsePeriodFromPathTest(unittest.TestCase):
    def test_month_name_with_year_in_filename(self):
        path = 'Jampea/2026/LAP. OPS. UPP JAMPEA PERIODE MEI 2026.pdf'
        self.assertEqual(parse_period_from_path(path), (2026, 5))

    def test_numeric_month_prefix(self):
        for bulan in (1, 2, 4):
            path = f'Jampea/2026/2026_{bulan:02d}_730104_JAMPEA.pdf'
            self.assertEqual(parse_period_from_path(path), (2026, bulan))

    def test_year_falls_back_to_parent_directory(self):
        path = 'Jampea/2026/LAPORAN OPERASIONAL JUNI.pdf'
        self.assertEqual(parse_period_from_path(path), (2026, 6))

    def test_unrecognised_month_is_not_defaulted_to_january(self):
        path = 'Jampea/2026/LAPORAN OPERASIONAL.pdf'
        self.assertEqual(parse_period_from_path(path), (2026, None))


class ParseIdNumberTest(unittest.TestCase):
    def test_indonesian_and_plain_formats(self):
        self.assertEqual(parse_id_number('2.097'), 2097.0)
        self.assertEqual(parse_id_number('1.150'), 1150.0)
        self.assertEqual(parse_id_number('1,5'), 1.5)
        self.assertEqual(parse_id_number('22,00'), 22.0)
        self.assertEqual(parse_id_number('21.63'), 21.63)
        self.assertEqual(parse_id_number('842'), 842.0)

    def test_non_numbers_are_none(self):
        for s in ('-', '', None, '######', 'SEMEN'):
            self.assertIsNone(parse_id_number(s))


def make_jampea_pdf():
    """A one-page report with 36 vertical rules (35 columns) and text placed per column."""
    edges = [100 + 24 * i for i in range(36)]
    doc = fitz.open()
    page = doc.new_page(width=1000, height=300)
    for x in edges:
        page.draw_line((x, 40), (x, 260))

    def put(y, **cols):
        for col, text in cols.items():
            page.insert_text((edges[int(col[1:])] + 1, y), text, fontsize=4)

    put(100, c0='1', c1='KAPALA', c2='RI', c4='21,63', c5='22', c7='01-06-26', c9='BULUKUMBA',
        c12='KLM', c13='02-06-26', c15='SELAYAR', c16='LPG', c17='27', c20='BERAS', c21='1,5',
        c28='3', c29='4')
    put(103, c16='SEMEN', c17='16')
    put(106, c17='296', c21='152')  # unlabelled section subtotal
    put(109, c1='JUMLAH', c17='999')
    put(112, c0='#', c1='KAPALB', c16='-', c17='-')
    return doc


class JampeaWordsPdfTest(unittest.TestCase):
    def test_vessels_split_and_subtotals_ignored(self):
        recs = parse_jampea_words_pdf(make_jampea_pdf(), 2026, 6)
        self.assertEqual([r['nama_kapal'] for r in recs], ['KAPALA', 'KAPALB'])
        a = recs[0]
        self.assertEqual((a['tahun'], a['bulan'], a['tipe_kapal']), (2026, 6, 'KLM'))
        self.assertEqual((a['gt'], a['loa']), (22.0, 21.63))
        self.assertEqual((a['penumpang_turun'], a['penumpang_naik']), (3, 4))
        self.assertEqual((a['barang_bongkar_ton'], a['barang_muat_ton']), (43.0, 1.5))
        self.assertEqual(a['jenis_barang_bongkar'], ['LPG', 'SEMEN'])
        self.assertEqual((a['pelabuhan_asal'], a['pelabuhan_tujuan']), ('BULUKUMBA', 'SELAYAR'))

    def test_page_without_ruled_table_gives_no_records(self):
        doc = fitz.open()
        doc.new_page().insert_text((50, 50), 'LAPORAN', fontsize=10)
        self.assertEqual(parse_jampea_words_pdf(doc, 2026, 6), [])


JUNE_PDF = os.path.join(RAW_DIR, 'Jampea/2026/LAPORAN OPERASIONAL JUNI.pdf')


@unittest.skipUnless(os.path.exists(JUNE_PDF), 'raw data not available')
class JampeaJuneRealFileTest(unittest.TestCase):
    def test_first_vessel_and_count(self):
        recs = parse_jampea_words_pdf(fitz.open(JUNE_PDF), 2026, 6)
        self.assertEqual(len(recs), 163)
        first = recs[0]
        self.assertEqual((first['nama_kapal'], first['tanggal_tiba']), ('REZKY AQILA', '01-06-26'))
        self.assertEqual((first['barang_bongkar_ton'], first['barang_muat_ton']), (27.0, 17.0))



MAY_2025_PDF = os.path.join(RAW_DIR, 'Jampea/2025/5. LAP. OPS. UPP JAMPEA PERIODE MEI 2025.pdf')


@unittest.skipUnless(os.path.exists(MAY_2025_PDF), 'raw data not available')
class JampeaTablePdfRealFileTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.recs = parse_jampea_pdf(fitz.open(MAY_2025_PDF), 2025, 5)

    def test_continuation_rows_are_counted(self):
        # SUMBER SURYA 02, arrived 03-05-25: unloads B.CAMPURAN 3 + BBM 2 + LPG 1, loads BERAS 4 + KOPRA 3 + LPG 1.
        rec = next(r for r in self.recs if r['nama_kapal'] == 'SUMBER SURYA 02' and r['tanggal_tiba'] == '03-05-25')
        self.assertEqual((rec['barang_bongkar_ton'], rec['barang_muat_ton']), (6.0, 8.0))
        self.assertEqual(rec['jenis_barang_bongkar'], ['B.CAMPURAN', 'BBM', 'LPG'])

    def test_column_number_row_is_not_a_vessel(self):
        self.assertEqual([r['nama_kapal'] for r in self.recs if not any(ch.isalpha() for ch in r['nama_kapal'])], [])

    def test_totals_match_report_grand_total(self):
        # JUMLAH TOTAL row of the report: unloaded 960 t, loaded 431 t, 0 animals unloaded, 388 loaded.
        # The report displays rounded sums; loaded cargo includes IKAN HIDUP 0,5, so the exact sum is 430.5.
        self.assertEqual(round(sum(r['barang_bongkar_ton'] for r in self.recs), 3), 960.0)
        self.assertEqual(round(sum(r['barang_muat_ton'] for r in self.recs), 3), 430.5)
        self.assertEqual(sum(r['hewan_turun_ekor'] for r in self.recs), 0)
        self.assertEqual(sum(r['hewan_naik_ekor'] for r in self.recs), 388)



JAN_2026_PDF = os.path.join(RAW_DIR, 'Jampea/2026/2026_01_730104_JAMPEA.pdf')


@unittest.skipUnless(os.path.exists(JAN_2026_PDF), 'raw data not available')
class JampeaPageBreakRealFileTest(unittest.TestCase):
    def test_cargo_rows_continue_across_page_break(self):
        # TIGA PUTRI BONERATE: B. CAMP 3, BBM 5, BERAS 3 on one page; ELPIJI 3, KURSI 1 after
        # JUMLAH DIPINDAHKAN / column numbers / JUMLAH PINDAHAN on the next.
        recs = parse_jampea_pdf(fitz.open(JAN_2026_PDF), 2026, 1)
        rec = next(r for r in recs if r['nama_kapal'].startswith('TIGA PUTRI'))
        self.assertEqual(rec['barang_bongkar_ton'], 15.0)
        self.assertEqual(rec['jenis_barang_bongkar'], ['B. CAMP', 'BBM', 'BERAS', 'ELPIJI', 'KURSI'])
        self.assertEqual(sum(r['barang_bongkar_ton'] for r in recs), 1082.0)


def jampea_row(**cols):
    """A 35-cell report row with the given cells set, e.g. jampea_row(c0='1', c1='KAPAL')."""
    row = [''] * 35
    for col, text in cols.items():
        row[int(col[1:])] = text
    return row


class JampeaRecordsFromRowsTest(unittest.TestCase):
    def test_sections_and_cargo_lines(self):
        rows = [
            jampea_row(c0='WILKER BONERATE'),
            jampea_row(c0='A', c1='PELAYARAN RAKYAT'),
            jampea_row(c0='1', c1='KAPALA', c16='SEMEN', c17='80', c20='KOPRA', c21='1,5', c22='SAPI', c23='7',
                       c30='GC', c31='103'),
            jampea_row(c16='BBM', c17='2', c22='KERBAU', c23='4'),
            jampea_row(c17='82', c21='1', c23='11', c28='0', c29='0'),
            jampea_row(c0='➔', c1='FERRY'),
            jampea_row(c0='1', c1='KAPALB', c16='-', c17='-', c20='NIHIL', c21='NIHIL'),
        ]
        a, b = jampea_records_from_rows(rows, 2025, 5)
        self.assertEqual((a['wilker'], a['jenis_pelayaran']), ('Wilker Bonerate', 'Pelayaran Rakyat'))
        self.assertEqual((b['wilker'], b['jenis_pelayaran']), ('Wilker Bonerate', 'Ferry'))
        self.assertEqual([(m['arah'], m['kategori'], m['komoditas_asli'], m['jumlah']) for m in a['muatan']], [
            ('bongkar', 'barang', 'SEMEN', 80.0), ('muat', 'barang', 'KOPRA', 1.5), ('muat', 'hewan', 'SAPI', 7.0),
            ('lanjutan', 'barang', 'GC', 103.0), ('bongkar', 'barang', 'BBM', 2.0), ('muat', 'hewan', 'KERBAU', 4.0)])
        # Through cargo is kept as a cargo line but not counted as unloaded or loaded.
        self.assertEqual((a['barang_bongkar_ton'], a['barang_muat_ton'], a['hewan_naik_ekor']), (82.0, 1.5, 11))
        self.assertEqual(b['muatan'], [])

    def test_quantity_with_unit_or_overflowing_type_text(self):
        # The type text can spill into the quantity cell ('PERALATAN PANCAN' + 'G 100'); units can be appended.
        rows = [jampea_row(c0='1', c1='KAPALA', c30='PERALATAN PANCAN', c31='G 100'),
                jampea_row(c31='25M3')]
        (a,) = jampea_records_from_rows(rows, 2025, 11)
        self.assertEqual([(m['komoditas_asli'], m['jumlah']) for m in a['muatan']],
                         [('PERALATAN PANCANG', 100.0), ('', 25.0)])

    def test_tonnage_entered_in_kilograms_is_converted(self):
        # A single cargo line of 238.080 'ton' is impossible at these ports; it is kilograms.
        rows = [jampea_row(c0='1', c1='KAPALA', c16='PIPA 508', c17='238.080'),
                jampea_row(c16='SEMEN', c17='40')]
        (a,) = jampea_records_from_rows(rows, 2025, 11)
        self.assertAlmostEqual(a['barang_bongkar_ton'], 278.08)
        self.assertEqual(a['muatan'][0]['catatan'], 'nilai laporan 238.080 dianggap kg')
        self.assertEqual(a['muatan'][1]['catatan'], '')


class JampeaSheetRowsTest(unittest.TestCase):
    def test_cells_become_report_text(self):
        ws = openpyxl.Workbook().active
        ws.append([1, 'HARAPAN KITA', 'RI', None, '28.75', 118, None, datetime.datetime(2026, 3, 3),
                   '05.30', 'TANJUNG PERAK'] + [None] * 6 + ['SEMEN', 14, '-', '-', 'KOPRA', 2.5]
                  + [None] * 13 + ['extra column'])
        (row,) = jampea_rows_from_sheet(ws)
        self.assertEqual(len(row), 35)
        self.assertEqual(row[:10], ['1', 'HARAPAN KITA', 'RI', '', '28.75', '118', '', '03-03-26', '05.30',
                                    'TANJUNG PERAK'])
        self.assertEqual(row[16:22], ['SEMEN', '14', '-', '-', 'KOPRA', '2.5'])
        self.assertEqual(parse_id_number(row[21]), 2.5)

    def test_thousands_read_as_decimals_are_restored(self):
        # The workbook was converted from the PDF, so '2.097' (GT 2097) is stored as the float 2.097.
        ws = openpyxl.Workbook().active
        ws.append([4, 'SABUK NUSANTARA-85', 'RI', 'PT.PELNI', '68.50', 2.097])
        (row,) = jampea_rows_from_sheet(ws)
        self.assertEqual(parse_id_number(row[5]), 2097.0)


MAR_2026_XLSX = os.path.join(RAW_DIR, 'Jampea/2026/2026_03_730104_JAMPEA.xlsx')


@unittest.skipUnless(os.path.exists(MAR_2026_XLSX), 'raw data not available')
class JampeaExcelRealFileTest(unittest.TestCase):
    def test_first_vessel_cargo(self):
        ws = openpyxl.load_workbook(MAR_2026_XLSX, read_only=True, data_only=True).worksheets[0]
        recs = jampea_records_from_rows(jampea_rows_from_sheet(ws), 2026, 3)
        rec = next(r for r in recs if r['nama_kapal'] == 'SUMBER SURYA 02')
        self.assertEqual(rec['tanggal_tiba'], '27-02-26')
        self.assertEqual((rec['barang_bongkar_ton'], rec['barang_muat_ton'], rec['hewan_naik_ekor']), (25.0, 10.0, 11))


class NormalizeKomoditasTest(unittest.TestCase):
    def test_spelling_variants_share_one_name(self):
        for raw in ('B.CAMP', 'B. Camp', 'B CAMPURAN', 'CAMPURAN', 'B.CAMPUR'):
            self.assertEqual(normalize_komoditas(raw), 'BARANG CAMPURAN')
        for raw in ('GC', 'G. C', 'KARGO UMUM', 'GC LANJUTAN'):
            self.assertEqual(normalize_komoditas(raw), 'GENERAL CARGO')
        self.assertEqual(normalize_komoditas('ELPJI KSG'), 'LPG KOSONG')
        self.assertEqual(normalize_komoditas('KEBAU'), 'KERBAU')
        self.assertEqual(normalize_komoditas("JAMBU MENTE'"), 'JAMBU MENTE')

    def test_unknown_names_pass_through_cleaned(self):
        self.assertEqual(normalize_komoditas('  rumput   laut kering '), 'RUMPUT LAUT KERING')
        self.assertEqual(normalize_komoditas(''), 'TIDAK DISEBUT')
        self.assertEqual(normalize_komoditas('0'), 'TIDAK DISEBUT')


class JampeaTotalDiffsTest(unittest.TestCase):
    def test_compares_with_grand_total_within_rounding(self):
        rows = [jampea_row(c0='1', c1='KAPALA', c16='SEMEN', c17='80', c20='IKAN HIDUP', c21='0,5', c30='GC', c31='10'),
                jampea_row(c0='JUMLAH TOTAL', c17='81', c19='0', c21='1', c23='0', c31='10')]
        recs = jampea_records_from_rows(rows, 2025, 5)
        self.assertEqual(jampea_total_diffs(recs, rows), [('barang_bongkar_ton', 80.0, 81.0)])

    def test_no_grand_total_row(self):
        self.assertIsNone(jampea_total_diffs([], [jampea_row(c0='1', c1='KAPALA')]))

if __name__ == '__main__':
    unittest.main()


# January 2024 as exported from the VHTS sheet TPK (Indonesian number format).
VHTS_JAN_2024 = ['2024', 'Januari', '39', '8495', '9335', '516', '336', '329', '22', '954', '12', '662', '8', '639',
                 '8495', '523', '6,157%', '6,81%', '674', '26', '977', '1.003', '1,49', '10,74', '1,87', '', '2,17', '1,48']


class HotelRecordTest(unittest.TestCase):
    def test_indicators_computed_from_base_counts(self):
        r = hotel_record_from_row(VHTS_JAN_2024)
        self.assertEqual((r['tahun'], r['bulan'], r['jumlah_sampel']), (2024, 1, 39))
        self.assertEqual((r['malam_kamar_tersedia'], r['malam_kamar_terpakai']), (8495, 523))
        self.assertEqual((r['tamu_asing'], r['tamu_indonesia'], r['tamu_total']), (12, 662, 674))
        self.assertEqual((r['malam_tamu_asing'], r['malam_tamu_indonesia'], r['malam_tamu']), (26, 977, 1003))
        self.assertEqual((r['tpk'], r['tptt']), (6.16, 10.74))
        self.assertEqual((r['rlmt_asing'], r['rlmt_indonesia'], r['rlmt_gabungan']), (2.17, 1.48, 1.49))

    def test_months_not_yet_filled_are_skipped(self):
        empty = ['2026', 'April'] + ['0'] * 14 + ['#DIV/0!', '', '0', '0', '0', '0'] + ['#DIV/0!'] * 3 + ['', '#DIV/0!', '#DIV/0!']
        self.assertIsNone(hotel_record_from_row(empty))
        self.assertIsNone(hotel_record_from_row(['', '', '', 'room_w', 'bed_w']))

    def test_no_foreign_guests_gives_no_average_stay(self):
        row = list(VHTS_JAN_2024)
        row[8] = row[10] = row[12] = '0'
        self.assertIsNone(hotel_record_from_row(row)['rlmt_asing'])

    def test_formula_cells_checked_against_computation(self):
        self.assertEqual(hotel_formula_diffs(hotel_record_from_row(VHTS_JAN_2024), VHTS_JAN_2024), [])
        row = list(VHTS_JAN_2024)
        row[16] = '7,5%'
        self.assertEqual(hotel_formula_diffs(hotel_record_from_row(row), row), [('tpk', 6.16, 7.5)])


@unittest.skipUnless(os.path.exists(VHTS_TPK_CSV), 'raw data not available')
class HotelRealFileTest(unittest.TestCase):
    def test_sheet_months_and_formulas(self):
        records, diffs = read_vhts_tpk(VHTS_TPK_CSV)
        self.assertEqual((records[0]['tahun'], records[0]['bulan']), (2024, 1))
        self.assertGreaterEqual(len(records), 27)
        self.assertEqual(diffs, [])
